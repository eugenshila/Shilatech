#!/usr/bin/env python3
"""Generate the Shilatech corporate overview and indicative online-cost deck.

Requires python-pptx (e.g. pip install python-pptx==1.0.2).
Run from repository root: python scripts/generate_corporate_presentation.py
The presentation describes capabilities in the checked-in application; cost estimates
come from the stakeholder brief and are not provider quotes.
"""
from pathlib import Path
from io import BytesIO
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'Shilatech_Corporate_Presentation.pptx'


def site_image(path):
    """Embed an existing website WebP asset without writing an intermediate file."""
    stream = BytesIO()
    Image.open(ROOT / path).convert('RGB').save(stream, format='PNG')
    stream.seek(0)
    return stream

LOGO = site_image('public/shilatech-logo-small.webp')
CATEGORY_ART = site_image('public/images/category-brakes.webp')
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

BG = '0B191C'; GREEN = '133B36'; GREEN2 = '1A4C42'; LIME = 'C5EB68'
CREAM = 'F4F5F0'; WHITE = 'FFFFFF'; INK = '172A2B'; MUTED = '667978'
PALE = 'E9EEE9'; RULE = 'D9E1DC'; AMBER = 'EFBA6A'; RED = '9E503F'


def rgb(h):
    return RGBColor.from_string(h)


def box(slide, x, y, w, h, fill=None, line=None, radius=False, lw=1):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        try: shape.adjustments[0] = 0.12
        except Exception: pass
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill or BG)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(lw)
    else:
        shape.line.fill.background()
    return shape


def txt(slide, text, x, y, w, h, size=16, color=INK, bold=False,
        align=PP_ALIGN.LEFT, font='Aptos', valign=MSO_ANCHOR.MIDDLE,
        margin=0, spacing=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = valign
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_after = Pt(spacing if spacing is not None else 0)
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = rgb(color)
    return shape


def line(slide, x, y, w, color=RULE, h=0.012):
    return box(slide, x, y, w, h, color)


def label(slide, text, x, y, w, dark=False):
    txt(slide, text.upper(), x, y, w, 0.24, 9, LIME if dark else GREEN2, True)


def slide_base(section, n, dark=False):
    s = prs.slides.add_slide(BLANK)
    box(s, 0, 0, 13.333, 7.5, BG if dark else CREAM)
    box(s, 0, 0, 0.11, 7.5, LIME if dark else GREEN2)
    LOGO.seek(0)
    s.shapes.add_picture(LOGO, Inches(0.63), Inches(0.17), width=Inches(0.92), height=Inches(0.59))
    txt(s, 'SHILATECH  /  AUTO SPARES', 1.67, 0.38, 4.0, 0.31, 12, WHITE if dark else GREEN, True)
    label(s, section, 9.15, 0.43, 3.55, dark)
    line(s, 0.62, 0.86, 12.03, GREEN2 if dark else RULE)
    line(s, 0.62, 7.08, 12.03, GREEN2 if dark else RULE)
    txt(s, 'CORPORATE PLATFORM OVERVIEW  •  OCT 2026', 0.62, 7.13, 6.8, 0.22, 9, '8FA39C' if dark else MUTED)
    txt(s, f'{n:02d} / 13', 11.85, 7.12, 0.8, 0.23, 9, LIME if dark else GREEN, True, PP_ALIGN.RIGHT)
    return s


def title(s, kicker, headline, sub=None, dark=False):
    label(s, kicker, 0.7, 1.12, 9, dark)
    txt(s, headline, 0.7, 1.42, 11.95, 0.61, 29, WHITE if dark else INK, True)
    if sub:
        txt(s, sub, 0.72, 2.02, 11.8, 0.36, 13, 'B5C6BF' if dark else MUTED)
    box(s, 0.72, 2.40, 0.66, 0.04, LIME)


def card(s, x, y, w, h, num, heading, desc, dark=False):
    fill = GREEN if dark else WHITE
    box(s, x, y, w, h, fill, GREEN2 if dark else RULE, True)
    txt(s, num, x + 0.24, y + 0.23, 0.58, 0.42, 21, LIME if dark else GREEN2, True)
    txt(s, heading, x + 0.24, y + 0.88, w - 0.48, 0.43, 18, WHITE if dark else INK, True)
    txt(s, desc, x + 0.24, y + 1.39, w - 0.48, h - 1.63, 12, 'CEDBD4' if dark else MUTED,
        valign=MSO_ANCHOR.TOP)


def notes(s, text):
    try:
        s.notes_slide.notes_text_frame.text = text
    except Exception:
        pass

# 01 — cover
s = slide_base('THE PLATFORM', 1, True)
box(s, 0.78, 1.58, 0.17, 4.42, LIME)
label(s, 'DIGITAL COMMERCE  /  OPERATIONS  /  GROWTH', 1.22, 1.57, 6.3, True)
txt(s, 'One platform.\nEvery part of\nthe operation.', 1.18, 2.02, 7.0, 2.82,
    42, WHITE, True, valign=MSO_ANCHOR.TOP)
txt(s, 'A connected digital foundation for Shilatech Auto Spares',
    1.22, 5.27, 6.45, 0.62, 17, 'C7D8CF')
box(s, 8.45, 1.63, 4.18, 4.66, GREEN, GREEN2, True)
# The site's existing category artwork is used as a visual, not as a product claim.
label(s, 'THE SHILATECH WEBSITE LOOK', 8.78, 1.94, 3.5, True)
CATEGORY_ART.seek(0)
s.shapes.add_picture(CATEGORY_ART, Inches(8.77), Inches(2.42), width=Inches(3.53), height=Inches(2.45))
box(s, 8.77, 5.03, 3.53, 0.82, '1B4B42', None, True)
txt(s, 'BLACK  /  GREEN  /  WHITE', 9.00, 5.26, 3.04, 0.30, 13, WHITE, True)
notes(s, 'Presentation based on the checked-in Shilatech application and documentation. This is a capability and deployment-planning overview; it does not claim that production deployment or live payment collection is complete.')

# 02 — proposition
s = slide_base('EXECUTIVE VIEW', 2)
title(s, 'THE OPPORTUNITY', 'Move from separate tasks to one connected workflow',
      'Make it easier to find parts, transact safely and coordinate the team behind every order.')
card(s, 0.70, 2.62, 3.83, 3.24, '01', 'Find the right part',
     'Catalogue search, vehicle brands, VIN lookup foundation and customer accounts give buyers a clearer path to purchase.')
card(s, 4.75, 2.62, 3.83, 3.24, '02', 'Sell with stock visibility',
     'Online checkout and the staff sales counter use shared inventory controls, reducing avoidable overselling.')
card(s, 8.80, 2.62, 3.83, 3.24, '03', 'Run the operation',
     'Warehouse, delivery, staff, finance and management views bring the day-to-day workflow into one system.')
box(s, 0.70, 6.20, 11.93, 0.48, PALE, None, True)
txt(s, 'POSITIONING', 0.96, 6.30, 1.08, 0.24, 10, GREEN, True)
txt(s, 'Custom-owned software  •  modular rollout  •  no recurring software licence',
    2.02, 6.28, 10.1, 0.28, 13, INK)

# 03 — customer journey
s = slide_base('CUSTOMER EXPERIENCE', 3)
title(s, 'FROM SEARCH TO SERVICE', 'A simpler path from vehicle to part',
      'The public website is designed around the questions a customer asks before and after checkout.')
journey = [
 ('01', 'DISCOVER', 'Browse brands, categories and featured parts.'),
 ('02', 'CHECK', 'Search by part or use VIN lookup as a fitment aid.'),
 ('03', 'CHOOSE', 'Review product details, cart and delivery options.'),
 ('04', 'ORDER', 'Submit checkout with the selected payment method.'),
 ('05', 'RETURN', 'Use account, workshop and support journeys.'),
]
for i, (num, head, desc) in enumerate(journey):
    x = 0.69 + i*2.43
    box(s, x, 2.76, 2.22, 2.84, WHITE, RULE, True)
    box(s, x, 2.76, 2.22, 0.08, GREEN2)
    txt(s, num, x+0.20, 3.05, 0.75, 0.46, 24, GREEN2, True)
    label(s, head, x+0.19, 3.70, 1.86)
    txt(s, desc, x+0.19, 4.10, 1.83, 1.18, 12, MUTED, valign=MSO_ANCHOR.TOP)
    if i < 4: txt(s, '→', x+2.19, 3.74, 0.28, 0.33, 18, GREEN2, True, PP_ALIGN.CENTER)
box(s, 0.69, 6.02, 11.94, 0.52, PALE, None, True)
txt(s, 'PAYMENT STATUS  •  Checkout methods are selectable; live gateway charging needs provider credentials and integration.',
    0.96, 6.13, 11.4, 0.27, 11, GREEN, True)
notes(s, 'VIN decoding and catalogue fitment filtering are a foundation, not a guarantee of compatibility. README.md says live gateway charges remain disabled until payment provider credentials are configured.')

# 04 — capability map
s = slide_base('PLATFORM MAP', 4)
title(s, 'ONE CONNECTED SYSTEM', 'Two experiences. One operational backbone.',
      'Customers see a focused storefront; staff see role-specific workspaces connected to the same data.')
for x, heading, tagline, items in [
    (0.70, 'CUSTOMER SIDE', 'Find, decide and request service',
     ['Product catalogue & brand filters', 'VIN lookup & My Garage', 'Cart, checkout & account', 'Workshop, contact & support']),
    (6.77, 'STAFF SIDE', 'Receive, sell and fulfil',
     ['Warehouse & barcode workflows', 'Counter POS & order handling', 'Delivery, garage & approvals', 'Finance, HR & management']),
]:
    box(s, x, 2.64, 5.86, 3.92, WHITE, RULE, True)
    box(s, x, 2.64, 5.86, 0.10, GREEN2)
    label(s, heading, x+0.30, 2.99, 4.9)
    txt(s, tagline, x+0.30, 3.37, 5.2, 0.39, 18, INK, True)
    for j, item in enumerate(items):
        yy = 4.00 + j*0.56
        box(s, x+0.31, yy+0.17, 0.12, 0.12, LIME)
        txt(s, item, x+0.61, yy, 4.89, 0.44, 14, INK)
box(s, 6.50, 4.30, 0.55, 0.55, LIME, None, True)
txt(s, '↔', 6.53, 4.38, 0.47, 0.31, 16, GREEN, True, PP_ALIGN.CENTER)

# 05 — stock discipline
s = slide_base('INVENTORY & SALES', 5, True)
title(s, 'THE OPERATIONAL CORE', 'Stock discipline across two sales channels',
      'Batch-based receiving and shared availability connect warehouse activity to online and counter sales.', True)
steps = [('RECEIVE', 'Log inbound batches\nand purchase costs'),
         ('CHECK', 'Reconcile sellable stock\nand online reservations'),
         ('SELL', 'Online order or\ncounter POS sale'),
         ('TRACE', 'Record movement,\nreceipt and audit trail')]
for i, (head, desc) in enumerate(steps):
    x = 0.70+i*3.02
    box(s, x, 2.77, 2.78, 2.18, GREEN, GREEN2, True)
    txt(s, f'0{i+1}', x+0.22, 2.97, 0.55, 0.43, 22, LIME, True)
    label(s, head, x+0.22, 3.51, 2.28, True)
    txt(s, desc, x+0.22, 3.89, 2.30, 0.79, 13, WHITE, valign=MSO_ANCHOR.TOP)
    if i<3: txt(s, '→', x+2.75, 3.56, 0.29, 0.34, 17, LIME, True, PP_ALIGN.CENTER)
box(s, 0.70, 5.37, 11.93, 1.13, '17342F', GREEN2, True)
txt(s, 'CURRENT BOUNDARY', 0.98, 5.59, 2.20, 0.30, 11, LIME, True)
txt(s, 'MAIN location is the active sales location; multi-branch transfers are a future phase, not a current capability.',
    3.13, 5.53, 9.08, 0.59, 14, WHITE)
notes(s, 'docs/WAREHOUSE_COUNTER_ROLLOUT.md describes shared batch stock, FIFO, reservations, MAIN location restrictions, idempotent counter sales and audit records. Electronic POS payments require independent verification before recording; no automatic gateway charging or verification.')

# 06 — fulfilment
s = slide_base('SERVICE OPERATIONS', 6)
title(s, 'AFTER THE SALE', 'Turn an order into a coordinated handoff',
      'Operational workspaces give teams a place to track the next action—not only the initial transaction.')
for x,y,num,head,body in [
    (0.70,2.62,'01','Order handling','Staff review order details and progress warehouse picking.'),
    (6.78,2.62,'02','Delivery','Delivery views help manage fulfilment and completion.'),
    (0.70,4.68,'03','Garage & workshop','Service requests and garage jobs sit alongside parts activity.'),
    (6.78,4.68,'04','Returns & exceptions','Returns and approvals create visible review points for exceptions.'),
]:
    box(s,x,y,5.85,1.78,WHITE,RULE,True)
    box(s,x+0.21,y+0.30,0.68,0.68,GREEN2,None,True)
    txt(s,num,x+0.29,y+0.46,0.50,0.26,13,LIME,True,PP_ALIGN.CENTER)
    txt(s,head,x+1.12,y+0.29,4.38,0.41,18,INK,True)
    txt(s,body,x+1.12,y+0.80,4.28,0.68,12,MUTED,valign=MSO_ANCHOR.TOP)
notes(s, 'Capability labels are drawn from pages and the website user manual. Do not present this as proof of an operational live deployment or integrated carrier booking.')

# 07 — finance & people
s = slide_base('BUSINESS MANAGEMENT', 7)
title(s, 'MANAGEMENT VIEW', 'Finance, people and approvals in one place',
      'Role-specific workspaces help managers follow the numbers and the decisions behind them.')
card(s, 0.70, 2.60, 3.83, 3.15, '01', 'Finance',
     'Receivables, recorded POS and invoice activity, and analysis of recorded costs and approved liabilities.')
card(s, 4.75, 2.60, 3.83, 3.15, '02', 'People',
     'HR records and a Kenya payroll preparation module with review and approval steps.')
card(s, 8.80, 2.60, 3.83, 3.15, '03', 'Governance',
     'Manager and administrator views, staff roles and approval checkpoints for controlled operations.')
box(s, 0.70, 6.03, 11.93, 0.53, PALE, None, True)
txt(s, 'BOUNDARY  •  Payroll is calculation and approval preparation; it does not send salaries or file statutory returns.',
    0.96, 6.14, 11.35, 0.27, 11, GREEN, True)
notes(s, 'docs/finance-analysis.md and docs/payroll-v1.md. Payroll v1 does not perform bank/M-Pesa transfers, remittances or payment-status recording. Finance analysis estimates require accountant review.')

# 08 — safeguards
s = slide_base('CONTROLS', 8, True)
title(s, 'BUILT FOR CONTROL', 'Operational safeguards worth preserving at launch',
      'A controlled go-live is as important as the software itself.', True)
for i, (head, body) in enumerate([
    ('Access by role', 'Staff routes and actions check the signed-in user and permissions.'),
    ('Stock integrity', 'Batch movement, reservations and discrepancy checks protect availability.'),
    ('Repeat-safe sales', 'Counter retries use request keys to avoid duplicate committed sales.'),
    ('Reviewable actions', 'Approval steps and audit records make important changes traceable.'),
]):
    x = 0.70 + (i%2)*6.08; y = 2.61 + (i//2)*1.74
    box(s,x,y,5.86,1.48,GREEN,GREEN2,True)
    txt(s, f'0{i+1}', x+0.20,y+0.22,0.58,0.34,17,LIME,True)
    txt(s,head,x+0.93,y+0.18,4.55,0.43,18,WHITE,True)
    txt(s,body,x+0.93,y+0.70,4.58,0.55,12,'CEDBD4',valign=MSO_ANCHOR.TOP)
box(s,0.70,6.28,11.93,0.42,'17342F',None,True)
txt(s,'GO-LIVE REQUIREMENT  •  Verify backup restore, payment handling and simultaneous sales on staging before opening.',
    0.96,6.37,11.38,0.23,11,LIME,True)

# 09 — deployment
s = slide_base('DEPLOYMENT BLUEPRINT', 9)
title(s, 'PUTTING IT ONLINE', 'A practical production architecture',
      'Proposed service layout for a managed, browser-accessible launch.')
layers = [
    (2.73,'01','CUSTOMER & STAFF','Web browsers  •  domain  •  free SSL'),
    (3.63,'02','APPLICATION','Next.js app  •  Vercel or Railway hosting'),
    (4.53,'03','DATA','Managed PostgreSQL  •  Supabase'),
    (5.43,'04','RESILIENCE','Email / alerts  •  encrypted off-site backup'),
]
for y,num,head,body in layers:
    box(s,0.70,y,11.93,0.73,WHITE,RULE,True)
    box(s,0.70,y,0.08,0.73,GREEN2)
    txt(s,num,0.95,y+0.18,0.57,0.36,16,GREEN2,True)
    txt(s,head,1.63,y+0.17,3.20,0.37,14,INK,True)
    txt(s,body,5.15,y+0.18,6.86,0.35,13,MUTED)
notes(s, 'Proposed architecture derived from README.md and stakeholder price assumptions. The email/alert setup and off-site backup bucket are deployment budget items, not a claim that these external services are provisioned today. SSL assumed free via host/certificate provider.')

# 10 — launch plan
s = slide_base('ROLLOUT PLAN', 10)
title(s, 'A CONTROLLED LAUNCH', 'Four gates before “go live”',
      'Sequence infrastructure, data validation and operational rehearsal instead of switching everything on at once.')
phases = [
    ('01','PREPARE','Confirm domains, provider accounts, roles and production secrets.'),
    ('02','STAGE','Migrate a copy, reconcile stock and test simultaneous workflows.'),
    ('03','REHEARSE','Restore a backup; run a controlled order, POS sale and delivery.'),
    ('04','LAUNCH','Deploy together, monitor transactions and keep a rollback plan.'),
]
for i,(num,head,body) in enumerate(phases):
    x=0.70+i*3.02
    box(s,x,2.78,2.78,3.43,WHITE,RULE,True)
    box(s,x,2.78,2.78,0.12,GREEN2)
    txt(s,num,x+0.23,3.13,0.85,0.58,28,GREEN2,True)
    label(s,head,x+0.23,4.04,2.3)
    txt(s,body,x+0.23,4.49,2.27,1.29,13,INK,valign=MSO_ANCHOR.TOP)
notes(s, 'See docs/WAREHOUSE_COUNTER_ROLLOUT.md for migration, reconciliation, backup and rollback ordering. Live gateway payment onboarding must be separately completed before charging customers; do not assume it is active.')

# 11 — exact cost matrix
s = slide_base('ONLINE COST  /  01', 11)
title(s, 'MONTHLY COMPONENTS', 'Cost of putting the system online',
      'Indicative monthly planning estimates in KES  •  based on stakeholder-supplied assumptions.')
xs=[0.70,4.13,6.95,9.77]; widths=[3.38,2.76,2.76,2.86]
for x,w,heading,fill,fc in zip(xs,widths,
        ['COST ITEM','TIER 1  /  PILOT','TIER 2  /  RECOMMENDED','TIER 3  /  VOLUME'],
        [GREEN,GREEN,GREEN2,GREEN],[WHITE,WHITE,LIME,WHITE]):
    box(s,x,2.49,w,0.51,fill)
    txt(s,heading,x+0.14,2.60,w-0.27,0.27,10,fc,True)
rows=[
 ('Domain + SSL','150–250','300','500'),
 ('Web app hosting','0–650','2,600','6,500–11,000'),
 ('Managed PostgreSQL','0–1,950','3,250','9,750–15,000'),
 ('Staff email + alerts','1,000–1,500','2,500–4,500','6,000–9,000'),
 ('Encrypted off-site backup','0','650','1,950'),
]
for j,(item,a,b,c) in enumerate(rows):
    y=3.00+j*0.54
    fill=WHITE if j%2==0 else PALE
    for x,w in zip(xs,widths): box(s,x,y,w,0.54,fill)
    txt(s,item,0.85,y+0.11,3.10,0.30,12,INK,j==0)
    for x,w,v in zip(xs[1:],widths[1:],[a,b,c]):
        txt(s,'KES '+v,x+0.12,y+0.10,w-0.24,0.33,13,GREEN2,True)
box(s,0.70,5.71,11.93,0.64,GREEN2)
txt(s,'MONTHLY TOTAL',0.85,5.87,3.03,0.29,12,WHITE,True)
for x,w,v in zip(xs[1:],widths[1:],['1,150–4,350','9,300–11,300','24,700–37,450']):
    txt(s,'KES '+v,x+0.12,5.83,w-0.24,0.35,14,LIME,True)
txt(s,'Totals are sums of the listed ranges; Tier 3 upper bound corrected from KES 38,450 to KES 37,450.',
    0.73,6.47,11.8,0.30,11,MUTED)
notes(s, 'Inputs supplied by the stakeholder. Tier 1 minima: 150+0+0+1000+0=1150; maxima: 250+650+1950+1500+0=4350. Tier 2 minima: 300+2600+3250+2500+650=9300; maxima: 300+2600+3250+4500+650=11300. Tier 3 minima: 500+6500+9750+6000+1950=24700; maxima: 500+11000+15000+9000+1950=37450. The supplied 38450 upper total did not match its components. Prices are not provider quotes; validate pricing, tax, FX, usage and availability before procurement.')

# 12 — yearly component matrix (12 × monthly values)
s = slide_base('ONLINE COST  /  02', 12)
title(s, 'YEARLY COMPONENTS', 'The annual cost, item by item',
      'Annualised from the monthly planning ranges: 12 × each listed component, in KES.')
xs=[0.70,4.13,6.95,9.77]; widths=[3.38,2.76,2.76,2.86]
for x,w,heading,fill,fc in zip(xs,widths,
        ['COST ITEM','TIER 1  /  PILOT','TIER 2  /  RECOMMENDED','TIER 3  /  VOLUME'],
        [GREEN,GREEN,GREEN2,GREEN],[WHITE,WHITE,LIME,WHITE]):
    box(s,x,2.49,w,0.51,fill)
    txt(s,heading,x+0.14,2.60,w-0.27,0.27,10,fc,True)
annual_rows=[
 ('Domain + SSL','1,800–3,000','3,600','6,000'),
 ('Web app hosting','0–7,800','31,200','78,000–132,000'),
 ('Managed PostgreSQL','0–23,400','39,000','117,000–180,000'),
 ('Staff email + alerts','12,000–18,000','30,000–54,000','72,000–108,000'),
 ('Encrypted off-site backup','0','7,800','23,400'),
]
for j,(item,a,b,c) in enumerate(annual_rows):
    y=3.00+j*0.54
    fill=WHITE if j%2==0 else PALE
    for x,w in zip(xs,widths): box(s,x,y,w,0.54,fill)
    txt(s,item,0.85,y+0.11,3.10,0.30,12,INK)
    for x,w,v in zip(xs[1:],widths[1:],[a,b,c]):
        txt(s,'KES '+v,x+0.12,y+0.10,w-0.24,0.33,12,GREEN2,True)
box(s,0.70,5.71,11.93,0.64,GREEN2)
txt(s,'YEARLY TOTAL',0.85,5.87,3.03,0.29,12,WHITE,True)
for x,w,v in zip(xs[1:],widths[1:],['13,800–52,200','111,600–135,600','296,400–449,400']):
    txt(s,'KES '+v,x+0.12,5.83,w-0.24,0.35,13,LIME,True)
txt(s,'Yearly totals use 12 months of service; Tier 3 upper bound follows the corrected monthly sum.',
    0.73,6.47,11.8,0.30,11,MUTED)
notes(s, 'Annual totals are the exact monthly sums times 12: Tier 1 KES 13,800–52,200; Tier 2 KES 111,600–135,600; Tier 3 KES 296,400–449,400. The stakeholder-supplied Tier 3 upper annual estimate KES 460,000 was based on the inconsistent upper monthly total. Domain Tier 2 KES 300/month annualises to KES 3,600, near the stakeholder approximation of KES 3,500/year. Provider prices, tax, USD/KES FX and usage must be checked at procurement.')

# 13 — decision
s = slide_base('ONLINE COST  /  03', 13, True)
label(s,'THE RECOMMENDATION',0.72,1.25,7.4,True)
txt(s,'Approve the corporate\nproduction budget.',0.70,1.64,8.60,1.49,36,WHITE,True,
    valign=MSO_ANCHOR.TOP)
box(s,0.72,3.51,7.54,1.35,GREEN,GREEN2,True)
txt(s,'KES 9,300–11,300',1.00,3.70,6.96,0.58,30,LIME,True)
txt(s,'per month  /  ~KES 112k–136k per year',1.02,4.31,6.87,0.28,14,WHITE)
box(s,8.61,1.47,3.97,4.58,GREEN,GREEN2,True)
label(s,'NO PLATFORM LICENCE',8.92,1.83,3.37,True)
txt(s,'KES 0 / mo',8.92,2.16,3.21,0.56,24,WHITE,True)
line(s,8.92,2.92,3.31,GREEN2)
label(s,'DARAJA STK PUSH ONBOARDING',8.92,3.21,3.4,True)
txt(s,'KES 0 setup  /  KES 0 monthly',8.92,3.57,3.23,0.72,17,WHITE,True)
txt(s,'Standard Paybill / Till transaction tariffs still apply.',8.92,4.47,3.25,1.02,12,'C9D8D1',valign=MSO_ANCHOR.TOP)
box(s,0.72,5.25,7.54,0.79,'17342F',None,True)
txt(s,'NEXT  •  Confirm quotes  →  Validate backup restore  →  Rehearse a controlled launch',
    0.98,5.45,7.01,0.39,12,WHITE,True)
txt(s,'Budget assumptions, not vendor quotes. USD items use ~KES 130 / USD; confirm FX, tax, usage limits and provider terms.\nPayment transaction charges and implementation work are excluded; live gateway charging is not yet enabled.',
    0.75,6.26,11.76,0.51,10,'B5C6BF',valign=MSO_ANCHOR.TOP)
notes(s, 'Budget is based on stakeholder-supplied estimates, not independently validated live provider quotations. The M-Pesa Daraja KES 0 setup/monthly claim is stakeholder-provided; independently verify onboarding eligibility and current Paybill/Till transaction tariff before contracting. Vercel $20≈KES2600, Supabase $25≈KES3250 and off-site bucket $5≈KES650 at an assumed 130 KES/USD; plans and usage may differ. Free SSL assumed. Checkout currently offers payment selection but live gateway charges remain disabled pending official credentials and integration (README.md).')

# Apply the website's approved black / grass-green / white palette to every slide.
# Recolor solid shapes and text separately so white cards become dark website panels,
# while white lettering remains white. All text and diagrams stay editable in PowerPoint.
panel_colors = {
    BG: '080A09', CREAM: '080A09', WHITE: '101610', PALE: '151915',
    GREEN: '101610', GREEN2: '243E18', LIME: '58B72A', RULE: '3B4935',
    '1B4B42': '151F15', '17342F': '151915',
}
text_colors = {
    BG: '080A09', INK: 'F6F7F5', WHITE: 'F6F7F5', MUTED: 'ABB5A7',
    GREEN: '9DDD71', GREEN2: '90DC5D', LIME: '9AE367',
    'CEDBD4': 'D8DDD7', 'B5C6BF': 'B8C6B7', '8FA39C': '9AA99B',
    'BED0C4': 'B8C6B7', 'DCE7E0': 'DCE7E0', 'C9D8D1': 'C9D8D1',
}
for slide in prs.slides:
    for shape in slide.shapes:
        try:
            current = str(shape.fill.fore_color.rgb)
            if current in panel_colors:
                shape.fill.fore_color.rgb = rgb(panel_colors[current])
        except (AttributeError, TypeError):
            pass  # picture and unfilled text boxes
        try:
            current = str(shape.line.color.rgb)
            if current in panel_colors:
                shape.line.color.rgb = rgb(panel_colors[current])
        except (AttributeError, TypeError):
            pass
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    try:
                        current = str(run.font.color.rgb)
                        if current in text_colors:
                            run.font.color.rgb = rgb(text_colors[current])
                    except (AttributeError, TypeError):
                        pass

prs.core_properties.title = 'Shilatech Auto Spares | Corporate Platform Overview'
prs.core_properties.subject = 'Platform capabilities, controlled launch and indicative monthly and yearly online costs'
prs.core_properties.author = 'Shilatech Auto Spares'
prs.core_properties.keywords = 'Shilatech, auto spares, Kenya, hosting, budget'
OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f'Wrote {OUT} ({len(prs.slides)} slides)')
